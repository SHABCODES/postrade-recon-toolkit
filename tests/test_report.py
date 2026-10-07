import openpyxl

from postrade_recon.report import build_report


def test_report_has_expected_sheets_and_live_formulas(result, tmp_path):
    path = build_report(result, "2026-10-07", tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)
    assert {"Summary", "Ageing", "All Breaks", "Trade Breaks", "Position Breaks", "Cash Breaks", "Margin Breaks", "Pricing Breaks", "Tolerances"} <= set(wb.sheetnames)
    assert wb["All Breaks"].max_row == len(result.breaks) + 1
    summary_formulas = [c.value for row in wb["Summary"].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("COUNTIFS" in f for f in summary_formulas)
    assert any("SUMIFS" in f for f in summary_formulas)
